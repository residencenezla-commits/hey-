"""Classeur Excel de la prime de rendement (SARL DPR AXXAM).

Reprend le document « Prime de rendement — affiche et fiches » :
objectif 10 wagons dépilés par équipe et par poste, 1 500 DA par wagon au-delà,
partagés à parts égales entre les présents, sous conditions de qualité.
Un classeur par mois : on saisit les postes et les présents, tout le reste se calcule.
"""
import datetime as dt
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SORTIE = sys.argv[1] if len(sys.argv) > 1 else "Prime_rendement_DPR_AXXAM.xlsx"
N_POSTES, N_PRES, N_EMP = 120, 1200, 60

BRIQUE = "8B2E16"
F = "Arial"
titre_f = Font(name=F, size=14, bold=True, color=BRIQUE)
ent_f = Font(name=F, size=10, bold=True, color="FFFFFF")
ent_fill = PatternFill("solid", fgColor=BRIQUE)
saisie_fill = PatternFill("solid", fgColor="FFF6D5")
saisie_f = Font(name=F, size=10, color="0000FF")
calc_fill = PatternFill("solid", fgColor="F2F2F2")
norm = Font(name=F, size=10)
gras = Font(name=F, size=10, bold=True)
note_f = Font(name=F, size=9, italic=True, color="595959")
fin = Side(style="thin", color="BFBFBF")
cadre = Border(left=fin, right=fin, top=fin, bottom=fin)
centre = Alignment(horizontal="center", vertical="center", wrap_text=True)
DA = '#,##0" DA";-#,##0" DA";"-"'
ENT = '#,##0;-#,##0;"-"'
PCT = '0.0%'
DATE = 'dd/mm/yyyy'

wb = Workbook()


def entete(ws, ligne, titres, largeurs):
    for i, (t, l) in enumerate(zip(titres, largeurs), 1):
        c = ws.cell(ligne, i, t)
        c.font, c.fill, c.alignment, c.border = ent_f, ent_fill, centre, cadre
        ws.column_dimensions[get_column_letter(i)].width = l
    ws.row_dimensions[ligne].height = 32


def cellule(ws, r, c, v=None, fmt=None, saisie=False, f=None):
    x = ws.cell(r, c, v)
    x.border = cadre
    x.font = f or (saisie_f if saisie else norm)
    if saisie:
        x.fill = saisie_fill
    elif isinstance(v, str) and v.startswith("="):
        x.fill = calc_fill
    if fmt:
        x.number_format = fmt
    return x


def titre(ws, texte, sous):
    ws["A1"] = "SARL DPR AXXAM — " + texte
    ws["A1"].font = titre_f
    ws["A2"] = sous
    ws["A2"].font = note_f


# ---------------------------------------------------------------- Paramètres
p = wb.active
p.title = "Paramètres"
titre(p, "Prime de rendement", "Cellules jaunes = à remplir. Cellules grises = calculées, ne pas toucher.")
p.column_dimensions["A"].width = 38
p.column_dimensions["B"].width = 18
p.column_dimensions["C"].width = 60
lignes = [
    ("Mois de la prime (1er jour du mois)", dt.date(2026, 10, 1), DATE, "Un classeur par mois : changer cette date pour un nouveau mois."),
    ("Objectif (wagons dépilés par équipe et par poste)", 10, ENT, "Document « Prime de rendement », décision de la direction."),
    ("Prime par wagon au-delà de l'objectif", 1500, DA, "Document « Prime de rendement », décision de la direction."),
    ("Seuil de casse et déclassés", 0.05, PCT, "À FIXER PAR LA DIRECTION (laissé en blanc dans le document) : 5 % est une valeur provisoire."),
    ("Paie du mois de", "", None, "Texte libre, ex. « novembre 2026 »."),
]
for i, (lib, val, fmt, note) in enumerate(lignes, 4):
    cellule(p, i, 1, lib, f=gras)
    cellule(p, i, 2, val, fmt, saisie=True)
    p.cell(i, 3, note).font = note_f
p["B7"].fill = PatternFill("solid", fgColor="FFFF00")
cellule(p, 10, 1, "Équipes", f=gras)
for i, e in enumerate(["Équipe A", "Équipe B", "Équipe C", "", ""], 11):
    cellule(p, i, 1, e, saisie=True)
p["C11"] = "Nommer les équipes comme à l'usine (A, B, C, matin, soir…)."
p["C11"].font = note_f
cellule(p, 17, 1, "Fin du mois (calculé)", f=gras)
cellule(p, 17, 2, "=EDATE(B4,1)", DATE)

regles = [
    "Comment remplir :",
    "1. Feuille « Employés » : la liste du personnel de production (une seule fois).",
    "2. Feuille « Postes » : une ligne par équipe et par jour, recopiée de la fiche journalière (wagons dépilés, casse, contrôles O/N).",
    "3. Feuille « Présences » : une ligne par personne présente ce jour-là dans cette équipe. Mettre O dans « Absence/retard injustifié » pour priver la personne de sa part.",
    "4. Les feuilles « Récap équipe », « État employés » et « Synthèse » se remplissent toutes seules.",
    "5. « État employés » donne le total à verser avec la paie, avec la colonne d'émargement.",
    "Règle : prime de l'équipe = (wagons dépilés − objectif) × prime par wagon, seulement si TOUS les contrôles sont à O et la casse sous le seuil.",
    "Part par personne = prime de l'équipe ÷ nombre de présents ayant droit, arrondie au dinar.",
]
for i, t in enumerate(regles, 20):
    p.cell(i, 1, t).font = gras if i == 20 else norm

EQ = "Paramètres!$A$11:$A$15"
dv_eq = lambda: DataValidation(type="list", formula1="=" + EQ, allow_blank=True)
def dv_on():
    return DataValidation(type="list", formula1='"O,N"', allow_blank=True)

# ---------------------------------------------------------------- Employés
e = wb.create_sheet("Employés")
titre(e, "Personnel de production", "Saisir une fois. Les noms servent de liste déroulante dans « Présences ».")
entete(e, 4, ["N°", "Nom et prénom", "Poste occupé", "Équipe habituelle"], [6, 32, 22, 18])
ex_emp = [("BENALI Karim", "Dépileur", "Équipe A"), ("HADDAD Samir", "Dépileur", "Équipe A"),
          ("MEZIANE Yacine", "Cariste", "Équipe A")]
dve = dv_eq(); e.add_data_validation(dve)
for k in range(N_EMP):
    r = 5 + k
    cellule(e, r, 1, k + 1, f=gras)
    nom, poste, eq = ex_emp[k] if k < len(ex_emp) else ("", "", "")
    cellule(e, r, 2, nom or None, saisie=True)
    cellule(e, r, 3, poste or None, saisie=True)
    cellule(e, r, 4, eq or None, saisie=True)
    dve.add(f"D{r}")
e["F5"] = "Lignes 1 à 3 = exemples (noms fictifs) : les remplacer par le vrai personnel."
e["F5"].font = note_f
e.freeze_panes = "A5"
NOMS = f"Employés!$B$5:$B${4 + N_EMP}"

# ---------------------------------------------------------------- Postes
s = wb.create_sheet("Postes")
titre(s, "Fiches journalières — une ligne par équipe et par poste",
      "Recopier la fiche du poste. Contrôles : O = oui, N = non. Une case vide = prime refusée.")
cols = ["Date", "Équipe", "Chef de poste", "Brique", "Poste (horaire)",
        "Wagons dépilés (1)", "Casse et déclassés (%)", "Paquets bien présentés et cerclés",
        "Sécurité respectée", "Poste propre", "Fiche visée chef de poste et production",
        "Conditions remplies", "Wagons au-delà (2)", "Prime de l'équipe (3)",
        "Présents ayant droit (4)", "Part par personne", "Motif si refus"]
entete(s, 4, cols, [11, 12, 18, 8, 12, 10, 11, 12, 10, 9, 13, 11, 10, 13, 10, 12, 26])
dvs = [dv_eq(), DataValidation(type="list", formula1='"B8,B12"', allow_blank=True), dv_on()]
for d in dvs: s.add_data_validation(d)
dv_pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
dv_pct.error = "Saisir un pourcentage, ex. 3 %"
dv_w = DataValidation(type="whole", operator="between", formula1="0", formula2="60", allow_blank=True)
s.add_data_validation(dv_pct); s.add_data_validation(dv_w)
exemples = [(dt.date(2026, 10, 1), "Équipe A", "Chef A", "B8", "6h-14h", 13, 0.03, "O", "O", "O", "O"),
            (dt.date(2026, 10, 1), "Équipe B", "Chef B", "B8", "14h-22h", 12, 0.07, "O", "O", "O", "O")]
for k in range(N_POSTES):
    r = 5 + k
    ex = exemples[k] if k < len(exemples) else (None,) * 11
    for c, v in enumerate(ex, 1):
        fmt = DATE if c == 1 else PCT if c == 7 else ENT if c == 6 else None
        cellule(s, r, c, v, fmt, saisie=True)
    for c in (8, 9, 10, 11): s.cell(r, c).alignment = centre
    cond = (f'=IF(OR(A{r}="",B{r}=""),"",IF(AND(F{r}<>"",G{r}<>"",G{r}<=Paramètres!$B$7,'
            f'H{r}="O",I{r}="O",J{r}="O",K{r}="O"),"O","N"))')
    cellule(s, r, 12, cond).alignment = centre
    cellule(s, r, 13, f'=IF(L{r}="","",MAX(0,F{r}-Paramètres!$B$5))', ENT)
    cellule(s, r, 14, f'=IF(L{r}="","",IF(L{r}="O",M{r}*Paramètres!$B$6,0))', DA)
    cellule(s, r, 15, f'=IF(L{r}="","",COUNTIFS(Présences!$A$5:$A${4+N_PRES},A{r},'
                      f'Présences!$B$5:$B${4+N_PRES},B{r},Présences!$D$5:$D${4+N_PRES},"<>O"))', ENT)
    cellule(s, r, 16, f'=IF(L{r}="","",IF(O{r}>0,ROUND(N{r}/O{r},0),0))', DA)
    cellule(s, r, 17, None, saisie=True)
    dvs[0].add(f"B{r}"); dvs[1].add(f"D{r}"); dv_pct.add(f"G{r}"); dv_w.add(f"F{r}")
    for c in "HIJK": dvs[2].add(f"{c}{r}")
s.conditional_formatting.add(f"L5:L{4+N_POSTES}", FormulaRule(formula=['L5="N"'], fill=PatternFill("solid", fgColor="F8CBAD")))
s.conditional_formatting.add(f"L5:L{4+N_POSTES}", FormulaRule(formula=['L5="O"'], fill=PatternFill("solid", fgColor="C6EFCE")))
s["A5"].comment = Comment("Lignes 5 et 6 = exemples (1er octobre) : à remplacer.", "DPR AXXAM")
s.cell(4, 15).comment = Comment("Compté automatiquement dans « Présences » : personnes de cette équipe ce jour-là, sans absence/retard injustifié.", "DPR AXXAM")
s.freeze_panes = "C5"

# ---------------------------------------------------------------- Présences
pr = wb.create_sheet("Présences")
titre(pr, "Présences par poste", "Une ligne par personne présente. La part se calcule toute seule.")
entete(pr, 4, ["Date", "Équipe", "Nom et prénom", "Absence/retard injustifié (O = pas de prime)", "Part (DA)"],
       [12, 14, 32, 18, 14])
dvp = [dv_eq(), DataValidation(type="list", formula1="=" + NOMS, allow_blank=True), dv_on()]
for d in dvp: pr.add_data_validation(d)
ex_pr = [(dt.date(2026, 10, 1), "Équipe A", "BENALI Karim", None),
         (dt.date(2026, 10, 1), "Équipe A", "HADDAD Samir", None),
         (dt.date(2026, 10, 1), "Équipe A", "MEZIANE Yacine", "O")]
P = f"Postes!$A$5:$A${4+N_POSTES}", f"Postes!$B$5:$B${4+N_POSTES}", f"Postes!$P$5:$P${4+N_POSTES}"
for k in range(N_PRES):
    r = 5 + k
    ex = ex_pr[k] if k < len(ex_pr) else (None,) * 4
    for c, v in enumerate(ex, 1):
        cellule(pr, r, c, v, DATE if c == 1 else None, saisie=True)
    pr.cell(r, 4).alignment = centre
    cellule(pr, r, 5, f'=IF(OR(A{r}="",B{r}="",C{r}=""),"",IF(D{r}="O",0,SUMIFS({P[2]},{P[0]},A{r},{P[1]},B{r})))', DA)
    dvp[0].add(f"B{r}"); dvp[1].add(f"C{r}"); dvp[2].add(f"D{r}")
pr["A5"].comment = Comment("Lignes 5 à 7 = exemples : à remplacer.", "DPR AXXAM")
pr.freeze_panes = "A5"

# ---------------------------------------------------------------- Récap équipe
rc = wb.create_sheet("Récap équipe")
titre(rc, "Récapitulatif mensuel de la prime — par équipe", "Choisir l'équipe en B3 (liste déroulante).")
rc["A3"] = "Équipe :"; rc["A3"].font = gras
cellule(rc, 3, 2, "Équipe A", saisie=True)
d = dv_eq(); rc.add_data_validation(d); d.add("B3")
rc["D3"] = "Mois :"; rc["D3"].font = gras
cellule(rc, 3, 5, "=Paramètres!B4", 'mmmm yyyy')
entete(rc, 5, ["Jour", "Date", "Wagons dépilés", "Au-delà de l'objectif", "Casse (%)", "Qualité (O/N)",
               "Prime équipe", "Présents", "Part / pers.", "Visa"], [6, 12, 10, 11, 9, 9, 14, 9, 12, 14])
plage = lambda c: f"Postes!${c}$5:${c}${4+N_POSTES}"
for j in range(1, 32):
    r = 5 + j
    cellule(rc, r, 1, j, f=gras)
    cellule(rc, r, 2, f'=IF(MONTH(Paramètres!$B$4+{j-1})=MONTH(Paramètres!$B$4),Paramètres!$B$4+{j-1},"")', DATE)
    n = f'COUNTIFS({plage("A")},B{r},{plage("B")},$B$3)'
    som = lambda c: f'=IF(OR(B{r}="",{n}=0),"",SUMIFS({plage(c)},{plage("A")},B{r},{plage("B")},$B$3))'
    cellule(rc, r, 3, som("F"), ENT)
    cellule(rc, r, 4, som("M"), ENT)
    cellule(rc, r, 5, f'=IF(OR(B{r}="",{n}=0),"",AVERAGEIFS({plage("G")},{plage("A")},B{r},{plage("B")},$B$3))', PCT)
    cellule(rc, r, 6, f'=IF(OR(B{r}="",{n}=0),"",IF(COUNTIFS({plage("A")},B{r},{plage("B")},$B$3,{plage("L")},"O")>0,"O","N"))').alignment = centre
    cellule(rc, r, 7, som("N"), DA)
    cellule(rc, r, 8, som("O"), ENT)
    cellule(rc, r, 9, som("P"), DA)
    cellule(rc, r, 10)
t = 37
cellule(rc, t, 1, "Total", f=gras); cellule(rc, t, 2)
for c, fmt in ((3, ENT), (4, ENT), (7, DA)):
    L = get_column_letter(c)
    cellule(rc, t, c, f"=SUM({L}6:{L}36)", fmt, f=gras)
for c in (5, 6, 8, 9, 10): cellule(rc, t, c)
rc["A39"] = "Jours travaillés :"; rc["A39"].font = gras
cellule(rc, 39, 3, '=COUNT(C6:C36)', ENT)
rc["D39"] = "Jours primés :"; rc["D39"].font = gras
cellule(rc, 39, 5, '=COUNTIF(G6:G36,">0")', ENT)
rc["F39"] = "Prime totale :"; rc["F39"].font = gras
cellule(rc, 39, 7, "=G37", DA, f=gras)
rc["A42"] = "Chef de production — nom et signature :"; rc["F42"] = "Direction — nom et signature :"
rc.freeze_panes = "A6"

# ---------------------------------------------------------------- État employés
et = wb.create_sheet("État employés")
titre(et, "État mensuel des primes — par employé", "À remettre à la paie. Se remplit depuis « Employés » et « Présences ».")
et["A3"] = "Mois :"; et["A3"].font = gras
cellule(et, 3, 2, "=Paramètres!B4", 'mmmm yyyy')
et["D3"] = "À verser avec la paie du mois de :"; et["D3"].font = gras
cellule(et, 3, 6, '=IF(Paramètres!B8="","",Paramètres!B8)')
entete(et, 5, ["N°", "Nom et prénom", "Poste", "Équipe", "Jours de présence", "Jours primés",
               "Total prime", "Émargement"], [6, 32, 18, 14, 11, 10, 15, 20])
PA = f"Présences!$A$5:$A${4+N_PRES}"
PC = f"Présences!$C$5:$C${4+N_PRES}"
PE = f"Présences!$E$5:$E${4+N_PRES}"
mois = '">="&Paramètres!$B$4', '"<"&Paramètres!$B$17'
for k in range(N_EMP):
    r, src = 6 + k, 5 + k
    cellule(et, r, 1, k + 1, f=gras)
    cellule(et, r, 2, f'=IF(Employés!B{src}="","",Employés!B{src})')
    cellule(et, r, 3, f'=IF(Employés!B{src}="","",IF(Employés!C{src}="","",Employés!C{src}))')
    cellule(et, r, 4, f'=IF(Employés!B{src}="","",IF(Employés!D{src}="","",Employés!D{src}))')
    crit = f'{PC},B{r},{PA},{mois[0]},{PA},{mois[1]}'
    cellule(et, r, 5, f'=IF(B{r}="","",COUNTIFS({crit}))', ENT)
    cellule(et, r, 6, f'=IF(B{r}="","",COUNTIFS({crit},{PE},">0"))', ENT)
    cellule(et, r, 7, f'=IF(B{r}="","",SUMIFS({PE},{crit}))', DA)
    cellule(et, r, 8)
t = 6 + N_EMP
cellule(et, t, 1, "Total général", f=gras)
for c in range(2, 9): cellule(et, t, c)
cellule(et, t, 7, f"=SUM(G6:G{t-1})", DA, f=gras)
et.cell(t + 2, 1, "Contrôle : total des parts saisies dans « Présences » pour le mois").font = note_f
cellule(et, t + 2, 7, f'=SUMIFS({PE},{PA},{mois[0]},{PA},{mois[1]})', DA)
et.cell(t + 3, 1, "(doit être égal au total général ; sinon un nom de « Présences » n'est pas dans « Employés »)").font = note_f
et.cell(t + 5, 1, "Chef de production :").font = gras
et.cell(t + 5, 4, "Comptabilité / Paie :").font = gras
et.cell(t + 5, 7, "Le Gérant :").font = gras
et.freeze_panes = "A6"

# ---------------------------------------------------------------- Synthèse
sy = wb.create_sheet("Synthèse", 0)
titre(sy, "Synthèse de la prime du mois", "Tableau de bord — tout est calculé.")
sy["A3"] = "Mois :"; sy["A3"].font = gras
cellule(sy, 3, 2, "=Paramètres!B4", 'mmmm yyyy')
entete(sy, 5, ["Équipe", "Postes travaillés", "Postes primés", "Wagons dépilés", "Moyenne wagons / poste",
               "Prime totale"], [16, 11, 11, 12, 13, 16])
for i in range(5):
    r = 6 + i
    eq = f"Paramètres!A{11+i}"
    crit = f'{plage("B")},A{r},{plage("A")},{mois[0]},{plage("A")},{mois[1]}'
    cellule(sy, r, 1, f'=IF({eq}="","",{eq})', f=gras)
    cellule(sy, r, 2, f'=IF(A{r}="","",COUNTIFS({crit}))', ENT)
    cellule(sy, r, 3, f'=IF(A{r}="","",COUNTIFS({crit},{plage("N")},">0"))', ENT)
    cellule(sy, r, 4, f'=IF(A{r}="","",SUMIFS({plage("F")},{crit}))', ENT)
    cellule(sy, r, 5, f'=IF(A{r}="","",IF(B{r}=0,0,D{r}/B{r}))', '0.0')
    cellule(sy, r, 6, f'=IF(A{r}="","",SUMIFS({plage("N")},{crit}))', DA)
cellule(sy, 11, 1, "Total", f=gras)
for c, fmt in ((2, ENT), (3, ENT), (4, ENT), (6, DA)):
    L = get_column_letter(c)
    cellule(sy, 11, c, f"=SUM({L}6:{L}10)", fmt, f=gras)
cellule(sy, 11, 5, '=IF(B11=0,0,D11/B11)', '0.0', f=gras)
sy["A13"] = "Objectif"; cellule(sy, 13, 2, "=Paramètres!B5", '0" wagons"')
sy["A14"] = "Prime par wagon au-delà"; cellule(sy, 14, 2, "=Paramètres!B6", DA)
sy["A15"] = "Seuil de casse"; cellule(sy, 15, 2, "=Paramètres!B7", PCT)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(SORTIE)
print("Écrit :", SORTIE)
