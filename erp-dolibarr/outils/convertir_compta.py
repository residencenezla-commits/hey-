#!/usr/bin/env python3
"""Convertit l'export DLG PC Compta (classeur « DLG_COMPTA… .xlsx » : INFO.DOSSIER,
JOURNAL, TAB_JRN, TAB_COM, TAB_AUX…) en fichiers CSV chargés par
init/30-comptabilite.php dans la comptabilité de Dolibarr.

Usage :
    python3 outils/convertir_compta.py DLG_COMPTA4_SARL_DPR_AXXAM_CONSOLIDE_2026.xlsx

Produit dans init/compta/ :
    plan_comptable.csv   comptes du plan (TAB_COM)
    journaux.csv         journaux et leur nature (TAB_JRN)
    auxiliaires.csv      comptes auxiliaires tiers (TAB_AUX)
    ecritures.csv        écritures du JOURNAL, regroupées par pièce
    controle.csv         totaux débit / crédit par compte, pour vérifier l'import

Règles :
- une pièce Dolibarr = un couple (journal, n° de pièce PC Compta) ; les lignes sans
  n° de pièce (réouverture) sont regroupées par journal et folio ;
- les écritures de réouverture datées du 31/12/2025 sont datées du 01/01/2026,
  premier jour de l'exercice, comme l'indique leur libellé ;
- le script refuse un fichier dont une pièce n'est pas équilibrée.
"""
import csv
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl

SORTIE = Path(__file__).resolve().parent.parent / "init" / "compta"
NATURES = {"Ventes": "vente", "Achats": "achat", "Trésorerie": "banque", "Stock": "stock",
           "Investissement": "divers", "Divers": "divers", None: "divers"}
# Journaux classés « Divers » dans PC Compta mais qui sont des achats ou de la trésorerie
CORRECTIONS = {"05": "achat", "15": "achat", "18": "banque", "24": "banque"}


def t(v):
    return "" if v is None else " ".join(str(v).split())


def n(v):
    return round(float(v or 0), 2)


def ecrire(nom, colonnes, lignes):
    with open(SORTIE / nom, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(colonnes)
        w.writerows(lignes)


def main(fichier):
    wb = openpyxl.load_workbook(fichier, read_only=True, data_only=True)
    SORTIE.mkdir(parents=True, exist_ok=True)
    feuille = lambda nom: [r for r in list(wb[nom].iter_rows(values_only=True))[1:] if r and r[0] not in (None, "")]

    dossier = {t(r[0]): t(r[1]) for r in wb["INFO.DOSSIER"].iter_rows(values_only=True) if r and r[0]}
    debut, fin = dossier.get("DEBUT_EXERCICE", "20260101"), dossier.get("FIN_EXERCICE", "20261231")

    comptes = OrderedDict((t(r[0]), t(r[1]) or t(r[0])) for r in feuille("TAB_COM"))
    ecrire("plan_comptable.csv", ["compte", "libelle", "classe"],
           [(c, l, c[0]) for c, l in sorted(comptes.items())])

    journaux = [(t(r[0]), t(r[1]), CORRECTIONS.get(t(r[0]), NATURES.get(r[7], "divers"))) for r in feuille("TAB_JRN")]
    ecrire("journaux.csv", ["code", "libelle", "nature"], journaux)

    aux = {t(r[0]): (t(r[1]), t(r[7]), t(r[8]), t(r[9]), t(r[10]), t(r[11])) for r in feuille("TAB_AUX")}
    ecrire("auxiliaires.csv", ["code", "libelle", "rue", "ville", "ai", "nif", "rc"],
           [(c,) + v for c, v in sorted(aux.items())])

    pieces = OrderedDict()
    for r in feuille("JOURNAL"):
        folio, ligne, piece, date, ref, libelle, jrn, compte, code_aux = (t(x) for x in r[:9])
        nat = t(r[10])
        if date < debut:
            date = debut  # réouverture au premier jour de l'exercice
        cle = (jrn, piece) if piece else (jrn, "F" + folio)
        pieces.setdefault(cle, []).append([jrn, piece, folio, ligne, date, ref, libelle, compte,
                                           comptes.get(compte, compte), code_aux,
                                           aux.get(code_aux, ("",))[0] if code_aux else "", nat,
                                           n(r[12]), n(r[13])])

    lignes, controle, desequilibres = [], defaultdict(lambda: [0.0, 0.0]), []
    for numero, (cle, lignes_piece) in enumerate(pieces.items(), start=1):
        if abs(sum(l[12] - l[13] for l in lignes_piece)) > 0.005:
            desequilibres.append(cle)
        for l in lignes_piece:
            lignes.append([numero] + l)
            controle[l[7]][0] += l[12]
            controle[l[7]][1] += l[13]
    if desequilibres:
        sys.exit(f"Pièces non équilibrées : {desequilibres[:10]}… Import annulé.")

    ecrire("ecritures.csv", ["piece_num", "journal", "piece", "folio", "ligne", "date", "reference", "libelle",
                             "compte", "libelle_compte", "aux", "libelle_aux", "nature", "debit", "credit"], lignes)
    ecrire("controle.csv", ["compte", "debit", "credit"],
           [(c, f"{d:.2f}", f"{cr:.2f}") for c, (d, cr) in sorted(controle.items())])
    ecrire("dossier.csv", ["cle", "valeur"], sorted(dossier.items()) + [("FICHIER_SOURCE", Path(fichier).name)])

    total_d = sum(l[13] for l in lignes)
    print(f"{dossier.get('NOM', '?')} — exercice {debut}–{fin}")
    print(f"{len(comptes)} comptes, {len(journaux)} journaux, {len(aux)} auxiliaires")
    print(f"{len(lignes)} lignes en {len(pieces)} pièces équilibrées, total débit = crédit = {total_d:,.2f} DA"
          .replace(",", " "))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
