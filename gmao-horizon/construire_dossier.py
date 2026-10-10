#!/usr/bin/env python3
"""Dossier de saisie HORIZON GMAO — atelier de préparation de l'argile, SARL DPR AXXAM.
Produit DPR-AXXAM_HORIZON_Preparation.xlsx (emplacements, équipements, plans préventifs, pièces).
Les équipements sont la liste type d'une ligne de préparation de briqueterie : à confirmer
et compléter avec le plan d'implantation des machines."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

A = 'Arial'
TITRE = Font(name=A, bold=True, size=14, color='7D2A17')
ENT = Font(name=A, bold=True, color='FFFFFF', size=10)
NORM = Font(name=A, size=10)
GRAS = Font(name=A, size=10, bold=True)
F_ENT = PatternFill('solid', fgColor='B4472B')
F_AFAIRE = PatternFill('solid', fgColor='FFFF00')     # à confirmer avec le plan
F_EX = PatternFill('solid', fgColor='F3E6D8')
T = Side(style='thin', color='BBBBBB')
CADRE = Border(left=T, right=T, top=T, bottom=T)

wb = Workbook()

def feuille(nom, titre, sous_titre, entetes, largeurs, lignes, a_confirmer=()):
    ws = wb.create_sheet(nom)
    ws['A1'] = titre; ws['A1'].font = TITRE
    ws['A2'] = sous_titre; ws['A2'].font = Font(name=A, italic=True, size=9, color='555555')
    for j, (h, w) in enumerate(zip(entetes, largeurs), 1):
        c = ws.cell(row=4, column=j, value=h); c.font = ENT; c.fill = F_ENT
        c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = CADRE
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, l in enumerate(lignes, 5):
        for j, v in enumerate(l, 1):
            c = ws.cell(row=i, column=j, value=v); c.font = NORM; c.border = CADRE
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if entetes[j - 1] in a_confirmer: c.fill = F_AFAIRE
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f"A4:{get_column_letter(len(entetes))}{4 + len(lignes)}"
    return ws

# ---------------------------------------------------------------- Mode d'emploi
ws = wb.active; ws.title = "Mode d'emploi"
lignes = [
    ("DPR AXXAM — Saisie dans HORIZON GMAO : atelier de préparation", TITRE),
    ("Briqueterie ZAC Helouane, Ighzer Amokrane (Béjaïa)", Font(name=A, italic=True, size=10)),
    ("", NORM),
    ("Ordre de saisie (HORIZON ne propose pas d'import de fichier, d'après le manuel v1.0 d'août 2026) :", GRAS),
    ("1. Feuille « 1 Emplacements » → Actifs > Emplacements > Nouvelle localisation, dans l'ordre des lignes (le parent d'abord).", NORM),
    ("2. Feuille « 2 Équipements » → Actifs & Équipements > Nouvel équipement : machines d'abord, puis leurs sous-ensembles (champ Équipement parent).", NORM),
    ("3. Imprimer les étiquettes : page Imprimer QR, une étiquette par Tag QR, collée sur la machine à l'endroit indiqué.", NORM),
    ("4. Feuille « 3 Plans préventifs » → Maintenance préventive > Créer un plan PM, avec la checklist et les consignes de sécurité.", NORM),
    ("5. Feuille « 4 Pièces » → Stock > Nouvelle pièce, avec équipement lié, stock initial et seuil d'alerte.", NORM),
    ("", NORM),
    ("Légende :", GRAS),
    ("Cellules jaunes = à confirmer ou compléter (plan d'implantation, plaques signalétiques, notices constructeur).", NORM),
    ("La liste des machines est la composition type d'une ligne de préparation de briqueterie (doseur → désagrégateur → laminoirs → mélangeur → stockage → reprise).", NORM),
    ("Elle sera ajustée au plan d'implantation des machines de l'usine dès réception : ajout, suppression ou renommage des lignes.", NORM),
    ("Fréquences et réglages : valeurs de départ prudentes ; les valider avec le responsable maintenance et les notices Marcheluzzo / constructeurs.", NORM),
    ("", NORM),
    ("Codification des Tags QR : DPR-<atelier>-<type>-<n°> ; sous-ensemble : tag de la machine + suffixe (-MOT moteur, -RED réducteur, -PAL paliers, -BND bande…).", NORM),
    ("Ateliers : PRE préparation · FAC façonnage · SEC séchoir · FOU four · CND conditionnement · UTI utilités (compresseur, eau, électricité).", NORM),
]
for i, (t, f) in enumerate(lignes, 1):
    ws.cell(row=i, column=1, value=t).font = f
ws.column_dimensions['A'].width = 140
ws['A12'].fill = F_AFAIRE

# ---------------------------------------------------------------- 1 Emplacements
EMP = [
    ("SITE", "Briqueterie DPR AXXAM — ZAC Helouane", "Site", "", "Usine de la ZAC Helouane, Ighzer Amokrane (Béjaïa)"),
    ("PRE", "Atelier Préparation de l'argile", "Bâtiment / atelier", "Briqueterie DPR AXXAM — ZAC Helouane", "De la réception de l'argile à l'alimentation de l'étireuse"),
    ("PRE-Z1", "Zone 1 — Réception et dosage", "Zone", "Atelier Préparation de l'argile", "Caissons doseurs alimentés à la chargeuse"),
    ("PRE-Z2", "Zone 2 — Broyage et laminage", "Zone", "Atelier Préparation de l'argile", "Désagrégateur, laminoirs dégrossisseur et finisseur"),
    ("PRE-Z3", "Zone 3 — Mélange et humidification", "Zone", "Atelier Préparation de l'argile", "Mélangeur, circuit d'eau"),
    ("PRE-Z4", "Zone 4 — Stockage / pourrissoir et reprise", "Zone", "Atelier Préparation de l'argile", "Stock de terre préparée, reprise vers le façonnage"),
    ("PRE-Z5", "Zone 5 — Convoyage", "Zone", "Atelier Préparation de l'argile", "Bandes de liaison entre machines"),
    ("PRE-EL", "Local électrique préparation", "Local technique", "Atelier Préparation de l'argile", "Armoires et variateurs de l'atelier"),
    ("FAC", "Atelier Façonnage (étireuse)", "Bâtiment / atelier", "Briqueterie DPR AXXAM — ZAC Helouane", "À détailler ensuite"),
    ("SEC", "Séchoir tunnel", "Bâtiment / atelier", "Briqueterie DPR AXXAM — ZAC Helouane", "À détailler ensuite"),
    ("FOU", "Four tunnel (84 m)", "Bâtiment / atelier", "Briqueterie DPR AXXAM — ZAC Helouane", "À détailler ensuite"),
    ("CND", "Déchargement, palettisation et expédition", "Bâtiment / atelier", "Briqueterie DPR AXXAM — ZAC Helouane", "À détailler ensuite"),
]
feuille("1 Emplacements", "1. Emplacements — Actifs > Emplacements", "À créer dans l'ordre : chaque parent existe avant ses enfants.",
        ["Code", "Nom de la localisation", "Type", "Localisation parente", "Description"], [10, 44, 20, 40, 60], EMP)

# ---------------------------------------------------------------- 2 Équipements
# (tag, nom, catégorie, emplacement, parent, fabricant, modèle, n° série, criticité, description)
M = []
def machine(tag, nom, cat, emp, crit, desc, sous=()):
    M.append((tag, nom, cat, emp, "", "", "", "", crit, desc))
    for suf, snom, scat, sdesc in sous:
        M.append((f"{tag}-{suf}", f"{nom} — {snom}", scat, emp, nom, "", "", "", crit, sdesc))
MOTRED = [("MOT", "moteur", "Moteur électrique", "Moteur d'entraînement (puissance et vitesse : plaque)"),
          ("RED", "réducteur", "Réducteur", "Réducteur / motoréducteur (type d'huile et quantité : plaque)")]
machine("DPR-PRE-DOS-01", "Caisson doseur n°1", "Doseur", "Zone 1 — Réception et dosage", "A",
        "Caisson d'alimentation à tablier (lattes ou bande) avec herse de dosage", MOTRED + [("TAB", "tablier / chaîne", "Transmission", "Tablier à lattes ou bande, chaînes, galets"), ("HER", "herse / arbre doseur", "Organe de travail", "Arbre à herse ou à palettes de dosage")])
machine("DPR-PRE-DOS-02", "Caisson doseur n°2 (ajout / dégraissant)", "Doseur", "Zone 1 — Réception et dosage", "B",
        "Second caisson si la ligne en comporte un (sable, argile 2)", MOTRED)
machine("DPR-PRE-DES-01", "Désagrégateur / émotteur", "Broyeur", "Zone 2 — Broyage et laminage", "A",
        "Broyage primaire des mottes et élimination des cailloux", MOTRED + [("CYL", "cylindres / rotor", "Organe de travail", "Cylindres dentés ou rotor à couteaux"), ("PAL", "paliers", "Paliers", "Paliers à roulements des arbres")])
machine("DPR-PRE-LAM-01", "Laminoir dégrossisseur", "Laminoir", "Zone 2 — Broyage et laminage", "A",
        "Laminoir à deux cylindres, écartement de dégrossissage", MOTRED + [("CYL", "cylindres", "Organe de travail", "Paire de cylindres, frettes, rectifieuse si présente"), ("PAL", "paliers", "Paliers", "Paliers des cylindres"), ("COU", "courroies", "Transmission", "Courroies trapézoïdales et poulies")])
machine("DPR-PRE-LAM-02", "Laminoir finisseur", "Laminoir", "Zone 2 — Broyage et laminage", "A",
        "Laminoir à deux cylindres, écartement fin (qualité de la pâte)", MOTRED + [("CYL", "cylindres", "Organe de travail", "Paire de cylindres, frettes"), ("PAL", "paliers", "Paliers", "Paliers des cylindres"), ("COU", "courroies", "Transmission", "Courroies trapézoïdales et poulies")])
machine("DPR-PRE-MEL-01", "Mélangeur humidificateur à double arbre", "Mélangeur", "Zone 3 — Mélange et humidification", "A",
        "Malaxage et réglage de l'humidité de la pâte", MOTRED + [("PAL", "pales et arbres", "Organe de travail", "Pales d'usure, bras, arbres"), ("ROU", "paliers", "Paliers", "Paliers des arbres"), ("EAU", "rampe d'eau", "Hydraulique", "Rampe, vannes et débitmètre d'eau")])
machine("DPR-PRE-EAU-01", "Pompe d'eau d'humidification", "Pompe", "Zone 3 — Mélange et humidification", "B",
        "Alimentation en eau du mélangeur", [("MOT", "moteur", "Moteur électrique", "Moteur de pompe")])
machine("DPR-PRE-STK-01", "Stockage / pourrissoir de terre préparée", "Stockage", "Zone 4 — Stockage / pourrissoir et reprise", "B",
        "Silo ou pourrissoir (repos de la pâte), avec chariot de remplissage si présent", [("CHA", "chariot / répartiteur", "Organe de travail", "Bande navette ou chariot de remplissage")])
machine("DPR-PRE-REP-01", "Excavateur / dispositif de reprise", "Reprise", "Zone 4 — Stockage / pourrissoir et reprise", "A",
        "Reprise de la terre stockée vers le façonnage (excavateur à godets, pont, ou chargeuse)", MOTRED + [("GOD", "godets / chaîne", "Organe de travail", "Chaîne à godets et dents")])
for n, (de, vers) in enumerate([("caisson doseur", "désagrégateur"), ("désagrégateur", "laminoir dégrossisseur"), ("laminoir dégrossisseur", "laminoir finisseur"),
                                ("laminoir finisseur", "mélangeur"), ("mélangeur", "stockage"), ("reprise", "façonnage")], 1):
    machine(f"DPR-PRE-CVY-{n:02d}", f"Convoyeur à bande n°{n} ({de} → {vers})", "Convoyeur", "Zone 5 — Convoyage", "B",
            f"Bande de liaison {de} → {vers}", [("MOT", "motoréducteur / tambour moteur", "Moteur électrique", "Entraînement de la bande"),
                                              ("BND", "bande et rouleaux", "Transmission", "Bande, rouleaux porteurs et de retour, racleurs")])
machine("DPR-PRE-ARM-01", "Armoire électrique préparation", "Électricité", "Local électrique préparation", "A",
        "Puissance, commande, variateurs et automate de la préparation")
machine("DPR-PRE-SEP-01", "Séparateur magnétique / détecteur de métaux", "Sécurité process", "Zone 5 — Convoyage", "B",
        "Protection des laminoirs contre les pièces métalliques (si installé)")
feuille("2 Équipements", "2. Équipements — Actifs & Équipements > Nouvel équipement",
        "Créer d'abord les machines (Équipement parent vide), puis leurs sous-ensembles. Criticité : A arrêt de la ligne · B gêne la production · C sans impact direct.",
        ["Tag QR Code", "Nom de l'équipement", "Catégorie", "Emplacement", "Équipement parent", "Fabricant", "Modèle", "N° de série", "Criticité", "Description"],
        [22, 50, 18, 36, 38, 16, 14, 14, 9, 52], M, a_confirmer=("Fabricant", "Modèle", "N° de série"))

# ---------------------------------------------------------------- 3 Plans préventifs
EPI = "Consignation électrique (cadenas) avant toute intervention ; EPI : casque, gants, chaussures, lunettes ; ne jamais intervenir sur une bande ou un cylindre en mouvement."
P = []
def plan(nom, equip, freq, duree, etapes, secu=EPI):
    P.append((nom, "Temps", freq, equip, duree, "\n".join(f"{i}. {e}" for i, e in enumerate(etapes, 1)), secu, "Technicien préparation"))
for eq in ("Caisson doseur n°1",):
    plan(f"PM Quotidien — {eq}", eq, "Quotidien", "15 min", ["Vérifier l'absence de blocs ou de corps étrangers dans le caisson", "Contrôler le tablier : lattes cassées, chaîne détendue (O/N)", "Écouter le réducteur : bruit ou vibration anormale (O/N)", "Nettoyer l'argile collée sous le tablier"])
for eq in ("Laminoir dégrossisseur", "Laminoir finisseur"):
    plan(f"PM Quotidien — {eq}", eq, "Quotidien", "20 min", ["Graisser les paliers des cylindres (pompe à graisse, nb de coups selon notice)", "Contrôler l'écartement des cylindres à la jauge — noter la valeur en mm", "Vérifier racleurs et nettoyage des cylindres", "Température des paliers au toucher / thermomètre infrarouge — noter °C"])
    plan(f"PM Hebdomadaire — {eq}", eq, "Hebdomadaire", "45 min", ["Tension et état des courroies (fissures, glissement)", "Contrôler l'usure des cylindres (ovalisation, frettes) — noter", "Serrage des boulons de bâti et des paliers", "Niveau d'huile du réducteur"])
    plan(f"PM Semestriel — {eq}", eq, "Tous les 6 mois", "1 jour", ["Mesurer le diamètre des cylindres en 3 points — noter", "Rectifier ou recharger les cylindres si l'usure dépasse la tolérance constructeur", "Remplacer les courroies si usées", "Vidange du réducteur (huile selon plaque)"])
plan("PM Quotidien — Désagrégateur / émotteur", "Désagrégateur / émotteur", "Quotidien", "15 min", ["Graisser les paliers", "Retirer pierres et corps étrangers", "Bruit / vibration anormale (O/N)"])
plan("PM Mensuel — Désagrégateur / émotteur", "Désagrégateur / émotteur", "Mensuel", "1 h", ["Contrôler l'usure des dents / couteaux — noter", "Serrage des fixations du rotor", "Niveau d'huile du réducteur"])
plan("PM Quotidien — Mélangeur humidificateur", "Mélangeur humidificateur à double arbre", "Quotidien", "20 min", ["Graisser les paliers des arbres", "Vérifier le débit d'eau et l'absence de buse bouchée", "Nettoyer la pâte durcie dans l'auge"])
plan("PM Mensuel — Mélangeur humidificateur", "Mélangeur humidificateur à double arbre", "Mensuel", "2 h", ["Mesurer l'usure des pales — remplacer celles hors tolérance", "Contrôler le jeu pales / auge", "Serrage des bras de pales", "Niveau d'huile du réducteur"])
for n in range(1, 7):
    plan(f"PM Hebdomadaire — Convoyeur n°{n}", f"Convoyeur à bande n°{n}", "Hebdomadaire", "20 min", ["Centrage de la bande et tension", "Rouleaux bloqués ou bruyants — noter le nombre", "État des racleurs et de la jonction de bande", "Arrêt d'urgence à câble : essai de fonctionnement"])
plan("PM Hebdomadaire — Excavateur / reprise", "Excavateur / dispositif de reprise", "Hebdomadaire", "1 h", ["Tension et graissage de la chaîne à godets", "Dents / godets usés ou déformés", "Fins de course et sécurités : essai"])
plan("PM Mensuel — Armoire électrique préparation", "Armoire électrique préparation", "Mensuel", "1 h", ["Dépoussiérage (aspiration, pas d'air comprimé sur les cartes)", "Filtres de ventilation : nettoyer ou remplacer", "Codes défauts des variateurs — relever", "Serrage des borniers de puissance (une fois par an, hors tension)"],
     "Intervention réservée à un électricien habilité ; armoire consignée ; vérifier l'absence de tension.")
plan("PM Trimestriel — Réducteurs de la préparation", "Laminoir finisseur", "Tous les 3 mois", "3 h", ["Relever le niveau et l'aspect de l'huile de chaque réducteur de la ligne", "Prélever un échantillon sur les réducteurs des laminoirs et du mélangeur", "Noter les heures de fonctionnement pour programmer les vidanges"])
feuille("3 Plans préventifs", "3. Plans préventifs — Maintenance préventive > Créer un plan PM",
        "Valeurs de départ prudentes : à valider avec le responsable maintenance et les notices constructeur. « Noter » = valeur à saisir dans la checklist.",
        ["Nom du plan", "Type de déclenchement", "Fréquence / intervalle", "Équipement", "Durée estimée", "Checklist (étapes)", "Instructions de sécurité", "Technicien responsable"],
        [42, 14, 16, 36, 11, 70, 50, 20], P, a_confirmer=("Technicien responsable",))

# ---------------------------------------------------------------- 4 Pièces
PI = [
    ("Bande transporteuse (largeur à relever)", "à compléter", "Convoyeur à bande n°1", "mètre", "", "", "1 longueur de rouleau entier"),
    ("Rouleau porteur de convoyeur", "à compléter", "Convoyeur à bande n°1", "pièce", "", "", "10"),
    ("Rouleau de retour de convoyeur", "à compléter", "Convoyeur à bande n°1", "pièce", "", "", "5"),
    ("Racleur de bande (lame)", "à compléter", "Convoyeur à bande n°1", "pièce", "", "", "4"),
    ("Roulement de palier — laminoirs", "réf. sur palier", "Laminoir finisseur", "pièce", "", "", "2"),
    ("Courroie trapézoïdale — laminoirs", "réf. sur courroie", "Laminoir finisseur", "pièce", "", "", "1 jeu complet"),
    ("Pale d'usure du mélangeur", "à compléter", "Mélangeur humidificateur à double arbre", "pièce", "", "", "1 jeu"),
    ("Dent / couteau de désagrégateur", "à compléter", "Désagrégateur / émotteur", "pièce", "", "", "1 jeu"),
    ("Latte de tablier du caisson doseur", "à compléter", "Caisson doseur n°1", "pièce", "", "", "5"),
    ("Huile réducteur (grade selon plaque)", "ex. ISO VG 220/320", "Laminoir finisseur", "litre", "", "", "1 vidange du plus gros réducteur"),
    ("Graisse paliers (EP2)", "à compléter", "Laminoir dégrossisseur", "kg", "", "", "1 seau"),
    ("Buse de rampe d'eau", "à compléter", "Mélangeur humidificateur à double arbre", "pièce", "", "", "4"),
]
feuille("4 Pièces", "4. Pièces de rechange critiques — Stock > Nouvelle pièce",
        "Coût unitaire en DZD, stock initial = inventaire physique, seuil d'alerte = consommation pendant le délai d'approvisionnement.",
        ["Nom", "Référence", "Équipement lié", "Unité", "Coût unitaire (DZD)", "Stock initial", "Seuil d'alerte conseillé"],
        [40, 20, 38, 9, 16, 12, 30], PI, a_confirmer=("Référence", "Coût unitaire (DZD)", "Stock initial"))

# Liste déroulante de criticité
ws = wb["2 Équipements"]
dv = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True); ws.add_data_validation(dv)
dv.add(f"I5:I{4 + len(M)}")

wb.save("DPR-AXXAM_HORIZON_Preparation.xlsx")
print(len(EMP), "emplacements,", len(M), "équipements,", len(P), "plans,", len(PI), "pièces")
